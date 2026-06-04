from pptx import Presentation
from pptx.util import Pt, Inches

def create_presentation():
    prs = Presentation()

    # Slide 1: Title Slide (SOFTWARE REQUIREMENT SPECIFICATION)
    slide_layout = prs.slide_layouts[0] # Title slide layout
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]

    title.text = "SOFTWARE\nREQUIREMENT\nSPECIFICATION\nfor\nRetail Store Management System iOS\nApplication"
    
    # Adjust title font size if needed
    for paragraph in title.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(28)

    subtitle.text = "Version History\nName\tDate\tVersion\nPrasad B S\t03/03/2026\t1.0"
    for paragraph in subtitle.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(18)

    # Slide 2: Team - who did what
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "1. Team & Responsibilities"
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Team Members & Roles:"
    
    p = tf.add_paragraph()
    p.text = "• Prasad B S - Project Lead / iOS Developer (Architecture & Core Systems)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Store Operations Team - UI/UX & Backend Integration"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Omnichannel Specialists - Payment & RFID Integration"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• After-Sales Experts - Service Workflow implementation"
    p.level = 1

    # Slide 3: What was customer ask?
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "2. The Customer Ask"
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Primary Objectives:"
    
    p = tf.add_paragraph()
    p.text = "• Build a Retail Store Management System (RSMS) tailored to a multi-country luxury goods retail chain."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Unify store operations, clienteling, product lifecycle, omnichannel fulfilment, and after-sales services."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Preserve the brand’s high-touch experience and compliance with global retail and payments regulations."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Provide a configurable product that can be adapted for various luxury retail formats, store sizes, and regional needs."
    p.level = 1

    # Slide 4: Team's understanding of requirement and approach
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "3. Team's Understanding & Approach"
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Understanding & Solution Approach:"
    
    p = tf.add_paragraph()
    p.text = "• Deep Dive into User Personas: Custom dashboards for Admin, Boutique Manager, Sales Associate, Inventory Controller, and After-Sales."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Technical Excellence: Built natively for iOS 26+ devices (iPhone/iPad) using SwiftUI, MVVM Architecture, and Swift Concurrency."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Omnichannel Focus: Centralized product master data, BOPIS/BORIS, Endless Aisle, and RFID-enabled stock management."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• High Reliability: Ensuring 0 memory leaks, 0 constraint issues, secure Passkey authentication, and offline resilience for POS."
    p.level = 1

    # Slide 5: Final ask team arrived at & Customer delight
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "4. Final Ask & Customer Delight"
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Exceeding Expectations:"
    
    p = tf.add_paragraph()
    p.text = "• Final Ask Re-defined: A highly scalable, secure (GDPR compliant), and performant ecosystem tailored for luxury retail."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Customer Delight Elements:"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "- AI-Assisted Clienteling: Using Core ML for intelligent cross-sell/up-sell recommendations."
    p.level = 2
    p = tf.add_paragraph()
    p.text = "- Advanced Vision Integration: Automated Intake & Diagnostics for After-Sales using camera & ML."
    p.level = 2
    p = tf.add_paragraph()
    p.text = "- Ultra-Smooth UI: Micro-animations and premium interface design ensuring a 'Wow' experience."
    p.level = 2

    # Slide 6: Final app (Video Demo, Live Demo, App on Phone)
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = "5. The Final App"
    
    content = slide.placeholders[1]
    tf = content.text_frame
    tf.text = "Demonstrations:"
    
    p = tf.add_paragraph()
    p.text = "• Video Demo: Highlighting the end-to-end luxury retail workflow."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• Live Demo: Showcasing real-time clienteling and RFID inventory scanning."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "• App on Phone: Hands-on experience with the RSMS application on target iOS devices."
    p.level = 1

    prs.save('RSMS_Presentation.pptx')
    print("Presentation created successfully as 'RSMS_Presentation.pptx'")

if __name__ == '__main__':
    create_presentation()
